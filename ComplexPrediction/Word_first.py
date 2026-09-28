import docx
from docx.shared import Pt

doc = docx.Document()

# style = doc.styles['Normal'] # Задаются параметры по умолчанию
# style.font.name = 'Arial'
# style.font.size = Pt(14)
#
# doc.add_paragraph("Абзац", 'Title') # Заголовок
# doc.add_paragraph("Абзац", 'List Bullet') # Ненумерованный список
#
# par1 = doc.add_paragraph('Первый абзац.')
# par2 = doc.add_paragraph('Второй абзац.')
# par3 = doc.add_paragraph('Третий абзац.')
#
# par1.add_run(' Дополнение первого абзаца. ').italic = True
# par2.add_run(' Дополнение второго абзаца. ').bold = True
# run = par3.add_run(' Дополнение третьего абзаца. ')
# run.bold = True
# run.underline = True

# doc.add_heading('Заголовок 0', 0)
# doc.add_heading('Заголовок 1', 1)
# doc.add_heading('Заголовок 2', 2)
# doc.add_heading('Заголовок 3', 3)
# doc.add_heading('Заголовок 4', 4)

# doc.add_paragraph('Это первая страница')
# # doc.add_page_break() # Разрыв страницы
# doc.add_section() # Разрыв раздела
# doc.add_paragraph('Это вторая страница')


# doc.add_picture('84291467_201954587649403_8431340612944278165_n.jpg', width=docx.shared.Cm(15)) # Вставка изображения

# table = doc.add_table(rows = 5, cols = 3)
#
# table.style = 'Table Grid'
#
# for row in range(5):
#     for col in range(3):
#         cell = table.cell(row, col)
#         cell.text = str(row)

doc.save('example.doc')
