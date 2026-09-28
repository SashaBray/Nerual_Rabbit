"""Ансамбль трёх моделей на честном датасете (климат-каналы + заполнение климатологией).

Датасет winter_wheat_honest, случайный сплит 70/15/15. Три модели обучаются на train:
  - свёрточная нейросеть (берём готовую winter_wheat_honest_e2e_nn/v12, обучена с аугментацией);
  - случайный лес (RandomForest);
  - линейная модель (Ridge).
Для леса и линейной вход — те же примеры при разных отсечках (аугментация климатологией),
сжатые по времени до L точек на канал и развёрнутые в вектор. Оценка — в день уборки (полные данные).
Метрики каждой модели по отдельности на test, затем СРЕДНЕЕ предсказание трёх (ансамбль), затем — на val.

Запуск:  python code/ensemble_honest.py
Выход:   workspace/reports/ensemble_honest/
"""

import json
import os

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch
from sklearn.ensemble import RandomForestRegressor
from sklearn.linear_model import Ridge

import config
import explore_models as E
import plots
import train_honest as TH
import train_nn as T

DATASET = 'winter_wheat_honest'
NN_VER = 'v12'
L = 50            # сжатие времени для леса/линейной (на канал)
SEED = 42
TREES = 200


def flat_at(rawf, climn, phn, a, b, n_years, size):
    X = TH.build_input(rawf, climn, phn, a, b, n_years, size, fill='clim', clim_channels=True)
    return E.to_vectors(E.downsample(X, L))            # (n, 23*L)


def main():
    outdir = os.path.join(config.REPORTS_DIR, 'ensemble_honest')
    os.makedirs(outdir, exist_ok=True)

    z, meta = TH.load_honest(DATASET)
    raw, clim, ph, valid = z['raw'], z['clim'], z['ph'], z['valid']
    ids, y = z['ids'], z['y'].astype(np.float32)
    keep = ~np.isnan(y) & (y >= 0) & (y <= E.YIELD_MAX)
    raw, clim, ph, valid, ids, y = raw[keep], clim[keep], ph[keep], valid[keep], ids[keep], y[keep]
    n_years = int(meta['duration_years']); size = int(meta['time_rows_size']); P = int(meta['n_params'])
    Tlen = n_years * size
    tr, te, va = TH.split_random(len(y), SEED)
    rawn, climn, phn, _ = TH.normalize(raw, clim, ph, tr)
    rawf = rawn.reshape(len(y), P, Tlen)
    harvest_abs = (n_years - 1) * size + TH.HARVEST_DOY
    print(f'Примеров {len(y)} (train {len(tr)}, test {len(te)}, val {len(va)}); '
          f'вход для леса/линейной: 23×{L}={23 * L} признаков')

    # --- обучающая матрица для леса/линейной: аугментация климатологией (сетка отсечек) ---
    cuts = [max(0, harvest_abs - d) for d in (0, 182, 365, 547, 730)]
    Xtr = np.vstack([flat_at(rawf[tr], climn[tr], phn[tr], a, min(a // size, n_years - 1), n_years, size)
                     for a in cuts])
    ytr = np.concatenate([y[tr]] * len(cuts))
    print(f'      обучающих строк (с аугментацией ×{len(cuts)} отсечек): {len(ytr)}')

    Xte = flat_at(rawf[te], climn[te], phn[te], harvest_abs, n_years - 1, n_years, size)
    Xva = flat_at(rawf[va], climn[va], phn[va], harvest_abs, n_years - 1, n_years, size)

    print('Обучение случайного леса…')
    rf = RandomForestRegressor(n_estimators=TREES, n_jobs=-1, random_state=SEED).fit(Xtr, ytr)
    print('Обучение линейной (Ridge)…')
    lin = Ridge(alpha=10.0).fit(Xtr, ytr)

    # --- свёрточная нейросеть v12 ---
    mdir = os.path.join(config.MODELS_DIR, 'winter_wheat_honest_e2e_nn', NN_VER)
    arch = json.load(open(os.path.join(mdir, 'architecture.json'), encoding='utf-8'))
    model = T.make_model(arch['arch_name'], arch['in_channels'], arch['in_length'])[0]
    model.load_state_dict(torch.load(os.path.join(mdir, 'model.pt'), map_location='cpu'))
    model.eval()

    def nn_pred(idx):
        X = TH.build_input(rawf[idx], climn[idx], phn[idx], harvest_abs, n_years - 1, n_years, size)
        return T.predict(model, X, 'cpu')

    preds = {}
    for split, idx, Xf in (('test', te, Xte), ('val', va, Xva)):
        p_nn, p_rf, p_lin = nn_pred(idx), rf.predict(Xf), lin.predict(Xf)
        preds[split] = {'y': y[idx], 'свёрточная': p_nn, 'лес': p_rf,
                        'линейная': p_lin, 'ансамбль (среднее)': (p_nn + p_rf + p_lin) / 3.0}

    rows = []
    for split in ('test', 'val'):
        d = preds[split]
        for name in ('свёрточная', 'лес', 'линейная', 'ансамбль (среднее)'):
            rows.append({'выборка': split, 'модель': name, **E.metrics(d['y'], d[name])})
    df = pd.DataFrame(rows)
    df.to_csv(os.path.join(outdir, 'ensemble_metrics.csv'), index=False, encoding='utf-8-sig')
    print('\n=== Метрики (прогноз в день уборки) ===')
    print(df[['выборка', 'модель', 'MSE', 'RMSE', 'R2', 'R_pearson']].round(3).to_string(index=False))

    _bar(df, outdir)
    _scatter(preds, outdir)
    _docx(df, outdir)
    print(f'\nГотово: {outdir}')


def _scatter(preds, outdir):
    """Точечный график реальная vs предсказанная урожайность ансамбля (test и val)."""
    def draw(lang):
        fig, axes = plt.subplots(1, 2, figsize=(12, 5.8))
        for ax, split in zip(axes, ('test', 'val')):
            d = preds[split]; yt = d['y']; yp = d['ансамбль (среднее)']
            lo = float(min(yt.min(), yp.min())); hi = float(max(yt.max(), yp.max()))
            ax.scatter(yt, yp, s=12, alpha=0.4, color='#d62728', edgecolors='none')
            ax.plot([lo, hi], [lo, hi], 'k--', lw=1, label='y = x')
            m = E.metrics(yt, yp)
            ax.set_xlabel(plots.tr('Реальная урожайность, ц/га', 'Actual yield, c/ha', lang))
            ax.set_ylabel(plots.tr('Предсказанная урожайность, ц/га', 'Predicted yield, c/ha', lang))
            ax.set_title(f'{split}: R²={m["R2"]:.3f}, RMSE={m["RMSE"]:.2f}, MSE={m["MSE"]:.1f}')
            ax.grid(True, alpha=0.3); ax.set_aspect('equal', 'box'); ax.legend(fontsize=8)
        fig.suptitle(plots.tr('Ансамбль: реальная и предсказанная урожайность — озимая пшеница',
                              'Ensemble: actual vs predicted yield — winter wheat', lang))
        fig.tight_layout()
        return fig
    return plots.bilingual(outdir, 'ensemble_scatter', draw)


def _bar(df, outdir):
    order = ['свёрточная', 'лес', 'линейная', 'ансамбль (среднее)']
    en = ['conv. NN', 'random forest', 'linear', 'ensemble (mean)']

    def draw(lang):
        fig, axes = plt.subplots(1, 2, figsize=(12, 5))
        labels = order if lang == 'ru' else en
        for ax, split in zip(axes, ('test', 'val')):
            sub = df[df['выборка'] == split].set_index('модель').loc[order]
            colors = ['#1f77b4', '#2ca02c', '#ff7f0e', '#d62728']
            ax.bar(range(len(order)), sub['MSE'].values, color=colors)
            ax.set_xticks(range(len(order)))
            ax.set_xticklabels(labels, rotation=20, ha='right', fontsize=8)
            ax.set_ylabel('MSE'); ax.set_title(f'MSE — {split}')
            for i, v in enumerate(sub['MSE'].values):
                ax.text(i, v, f'{v:.1f}', ha='center', va='bottom', fontsize=8)
            ax.grid(True, axis='y', alpha=0.3)
        fig.suptitle(plots.tr('Ансамбль трёх моделей — озимая пшеница (честный датасет)',
                              'Three-model ensemble — winter wheat (honest dataset)', lang))
        fig.tight_layout()
        return fig
    return plots.bilingual(outdir, 'ensemble_mse', draw)


def _docx(df, outdir):
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading('Ансамбль трёх моделей — озимая пшеница (честный датасет)', 0)
    doc.add_paragraph(
        'Датасет winter_wheat_honest (климат-каналы + заполнение будущего средней климатологией района), '
        'случайный сплит 70/15/15. Три модели обучены на train: свёрточная нейросеть (v12, с '
        'аугментацией), случайный лес и линейная (Ridge); для леса и линейной время сжато до '
        f'{L} точек на канал. Оценка — в день уборки (полные данные). Ансамбль = среднее предсказание '
        'трёх моделей.')
    for split in ('test', 'val'):
        doc.add_heading(f'Метрики — {split}', level=1)
        sub = df[df['выборка'] == split][['модель', 'MSE', 'RMSE', 'R2', 'R_pearson']].round(3)
        E._table_from_df(doc, sub)
    png = os.path.join(outdir, 'ensemble_mse_ru.png')
    if os.path.exists(png):
        doc.add_heading('График MSE по моделям', level=1)
        doc.add_picture(png, width=Inches(6.3))
    sc = os.path.join(outdir, 'ensemble_scatter_ru.png')
    if os.path.exists(sc):
        doc.add_heading('Реальная и предсказанная урожайность (ансамбль)', level=1)
        doc.add_picture(sc, width=Inches(6.3))
    out = os.path.join(outdir, 'report_ensemble_honest.docx')
    try:
        doc.save(out)
    except PermissionError:
        out = os.path.join(outdir, 'report_ensemble_honest_new.docx'); doc.save(out)
    print(f'Отчёт: {out}')


if __name__ == '__main__':
    main()
