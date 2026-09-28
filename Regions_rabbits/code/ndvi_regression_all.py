"""Классическая регрессия по NDVI для ВСЕХ культур из БД.

Тот же метод, что и для озимой пшеницы: по каждому району экспонента урожай = a·exp(b·maxNDVI),
где maxNDVI — максимум NDVI ДО УБОРКИ (дни 0..harvest_doy целевого года, из harvest_dates.csv).
Каждая культура — отдельная глава; в конце большая сводная таблица по всем культурам.

Данные берутся из готовых датасетов workspace/datasets/<культура>/ (matrix.csv: NDVI = признак 0,
блок целевого года; scalar.csv: id_region, id_district, year, productive). Сплит train/test/val 70/15/15.

Запуск:  python code/ndvi_regression_all.py
Выход:   workspace/reports/ndvi_all/ (+ docs/Регрессия_NDVI_по_культурам.docx)
"""
import json
import os

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

import config
import explore_models as E
import plots
import train_honest as TH
from ndvi_regression import fit_exp, pred_exp, pearson, MIN_FIT, SEED

NAMES = [
    ('winter_wheat_all', 'Озимая пшеница'), ('spring_wheat_all', 'Яровая пшеница'),
    ('barley_winter_all', 'Ячмень озимый'), ('barley_spring_all', 'Ячмень яровой'),
    ('rye_winter_all', 'Рожь озимая'), ('rye_spring_all', 'Рожь яровая'),
    ('oats_all', 'Овёс'), ('corn_all', 'Кукуруза'), ('soy_all', 'Соя'),
    ('sunflower_all', 'Подсолнечник'), ('buckwheat_all', 'Гречиха'), ('peas_all', 'Горох'),
    ('flax_all', 'Лён'), ('sugarbeet_all', 'Сахарная свёкла'), ('potato_all', 'Картофель'),
    ('rice_all', 'Рис'),
]


def harvest_map():
    df = pd.read_csv(config.HARVEST_DATES_FILE if hasattr(config, 'HARVEST_DATES_FILE')
                     else os.path.join(config.CONFIG_DIR, 'harvest_dates.csv'))
    return dict(zip(df['task'], df['harvest_doy'].astype(int)))


def load_maxndvi(name, hdoy, size=365):
    base = os.path.join(config.DATASETS_DIR, name)
    c0 = 2 * size                                    # признак 0 (ndvi), целевой блок (3-й год)
    cols = list(range(c0, c0 + int(hdoy)))           # дни 0..harvest_doy до уборки
    mat = pd.read_csv(os.path.join(base, 'matrix.csv'), sep=' ', header=None,
                      usecols=cols, dtype='float32').values
    sc = pd.read_csv(os.path.join(base, 'scalar.csv'), sep=' ', header=None).values
    return (sc[:, 0].astype(int), sc[:, 1].astype(int), sc[:, 3].astype(float), mat.max(axis=1))


def run_culture(name, hdoy):
    reg, dist, y, mx = load_maxndvi(name, hdoy)
    keep = ~np.isnan(y) & (y >= 0) & (y <= E.YIELD_MAX) & (mx > 0)
    reg, dist, y, mx = reg[keep], dist[keep], y[keep], mx[keep]
    n = len(y)
    terr = reg * 100000 + dist
    tr, te, va = TH.split_random(n, SEED)
    in_tr = np.zeros(n, bool); in_tr[tr] = True
    glob = float(y[tr].mean())
    ycap = float(min(E.YIELD_MAX, 1.5 * y[tr].max()))    # культурный потолок (защита от выбросов)
    pred = np.full(n, np.nan); meth = np.zeros(n, bool)
    corrs, drows = [], []
    for t in set(terr.tolist()):
        sel = terr == t
        tr_m = sel & in_tr
        xtr = mx[tr_m]
        fit = fit_exp(xtr, y[tr_m])
        dmean = float(y[tr_m].mean()) if tr_m.sum() else glob
        xlo, xhi = (float(xtr.min()), float(xtr.max())) if tr_m.sum() else (0.0, 1.0)
        for i in np.where(sel)[0]:
            if fit:
                xc = min(max(float(mx[i]), xlo), xhi)    # без экстраполяции экспоненты за пределы данных
                pred[i] = min(pred_exp(fit, xc), ycap)
                meth[i] = True
            else:
                pred[i] = dmean
        r = pearson(mx[sel], y[sel])
        ridx = np.where(sel)[0][0]
        corrs.append(r)
        drows.append({'id_region': int(reg[ridx]), 'id_district': int(dist[ridx]),
                      'n': int(sel.sum()), 'corr_ndvi_yield': round(r, 3) if r == r else np.nan,
                      'exp_a': round(np.exp(fit[0]), 2) if fit else np.nan,
                      'exp_b': round(fit[1], 3) if fit else np.nan})
    cc = np.array([c for c in corrs if c == c])
    mte, mva = E.metrics(y[te], pred[te]), E.metrics(y[va], pred[va])
    summ = {'культура': None, 'n': n, 'n_районов': len(set(terr.tolist())),
            'test_R2': round(mte['R2'], 3), 'test_RMSE': round(mte['RMSE'], 2), 'test_MSE': round(mte['MSE'], 1),
            'val_R2': round(mva['R2'], 3), 'val_RMSE': round(mva['RMSE'], 2), 'val_MSE': round(mva['MSE'], 1),
            'медиана_corr': round(float(np.median(cc)), 2) if len(cc) else np.nan,
            'доля_|r|>0.5': round(float((np.abs(cc) > 0.5).mean()), 2) if len(cc) else np.nan,
            'доля_exp': round(float(meth[te].mean()), 2)}
    return summ, pd.DataFrame(drows), (y[te], pred[te])


def main():
    outdir = os.path.join(config.REPORTS_DIR, 'ndvi_all')
    os.makedirs(outdir, exist_ok=True)
    hmap = harvest_map()
    rows, scatters = [], {}
    for name, rus in NAMES:
        if not os.path.exists(os.path.join(config.DATASETS_DIR, name, 'matrix.csv')):
            print(f'{rus}: нет датасета — пропуск'); continue
        hd = hmap.get(name, 210)
        print(f'{rus} ({name}, harvest_doy={hd})…', flush=True)
        summ, ddf, sc = run_culture(name, hd)
        summ['культура'] = rus
        rows.append(summ); scatters[rus] = sc
        ddf.to_csv(os.path.join(outdir, f'district_{name}.csv'), index=False, encoding='utf-8-sig')
        print(f'   test R²={summ["test_R2"]} MSE={summ["test_MSE"]} | val R²={summ["val_R2"]} | '
              f'медиана corr={summ["медиана_corr"]}', flush=True)
    summary = pd.DataFrame(rows).sort_values('test_R2', ascending=False).reset_index(drop=True)
    summary.to_csv(os.path.join(outdir, 'summary_all_cultures.csv'), index=False, encoding='utf-8-sig')
    print('\n=== СВОДКА ===')
    print(summary[['культура', 'n', 'test_R2', 'test_MSE', 'val_R2', 'медиана_corr', 'доля_|r|>0.5']].to_string(index=False))

    bar = _summary_bar(summary, outdir)
    sc_imgs = _scatters(scatters, outdir)
    _docx(outdir, summary, scatters, bar, sc_imgs)
    print(f'\nГотово: {outdir}')


def _summary_bar(summary, outdir):
    def draw(lang):
        fig, ax = plt.subplots(figsize=(12, 6))
        s = summary.sort_values('test_R2')
        x = np.arange(len(s))
        ax.barh(x, s['test_R2'], color='#1f77b4', label=plots.tr('R² на тесте', 'test R²', lang))
        ax.plot(s['медиана_corr'], x, 'o', color='#d62728',
                label=plots.tr('медиана corr(NDVI,урожай)', 'median corr(NDVI,yield)', lang))
        ax.set_yticks(x); ax.set_yticklabels(s['культура'], fontsize=8)
        ax.set_xlabel(plots.tr('R² (тест) / медианная корреляция', 'R² (test) / median correlation', lang))
        ax.set_title(plots.tr('Классическая регрессия по NDVI — по культурам',
                              'Classical NDVI regression — by culture', lang))
        ax.legend(); ax.grid(True, axis='x', alpha=0.3); fig.tight_layout(); return fig
    return plots.bilingual(outdir, 'summary_cultures', draw)


def _scatters(scatters, outdir):
    imgs = {}
    for rus, (yt, pr) in scatters.items():
        slug = ''.join(ch if ch.isalnum() else '_' for ch in rus)

        def draw(lang, yt=yt, pr=pr, rus=rus):
            fig, ax = plt.subplots(figsize=(5.5, 5.5))
            lo, hi = 0, float(max(yt.max(), pr.max()) + 1)
            ax.scatter(yt, pr, s=10, alpha=0.4, color='#2ca02c', edgecolors='none')
            ax.plot([lo, hi], [lo, hi], 'k--', lw=1)
            m = E.metrics(yt, pr)
            ax.set_title(f'{rus}: R²={m["R2"]:.3f}, RMSE={m["RMSE"]:.2f}')
            ax.set_xlabel(plots.tr('Реальная, ц/га', 'Actual, c/ha', lang))
            ax.set_ylabel(plots.tr('Прогноз по NDVI, ц/га', 'NDVI prediction, c/ha', lang))
            ax.grid(True, alpha=0.3); ax.set_aspect('equal', 'box'); fig.tight_layout(); return fig
        imgs[rus] = plots.bilingual(outdir, f'scatter_{slug}', draw)
    return imgs


def _docx(outdir, summary, scatters, bar, sc_imgs):
    from docx import Document
    from docx.shared import Inches
    docs = os.path.join(os.path.dirname(os.path.dirname(config.REPORTS_DIR)), 'docs')
    os.makedirs(docs, exist_ok=True)
    doc = Document()
    doc.add_heading('Классическая регрессия по NDVI — все культуры', 0)
    doc.add_paragraph(
        'Метод (как для озимой пшеницы) применён ко всем культурам из БД. Для каждого РАЙОНА по всем '
        'годам строится связь «максимум NDVI до уборки → урожайность» и аппроксимируется ЭКСПОНЕНТОЙ '
        'урожай = a·exp(b·maxNDVI). Максимум NDVI берётся по дням 0..день_уборки целевого года '
        '(harvest_dates.csv) — для озимых это весенний пик, для яровых/поздних — летний, до уборки. '
        f'Сплит train/test/val = 70/15/15; экспонента подгоняется на train-точках района (мин. {MIN_FIT}), '
        'иначе откат к средней урожайности района. Прогноз НЕ экстраполируется за пределы наблюдённого '
        'диапазона NDVI района (иначе экспонента «взрывается») и ограничен культурным потолком урожайности.')

    doc.add_heading('Сводная таблица по культурам (сортировка по R² на тесте)', level=1)
    cols = ['культура', 'n', 'n_районов', 'test_R2', 'test_RMSE', 'test_MSE',
            'val_R2', 'val_RMSE', 'медиана_corr', 'доля_|r|>0.5', 'доля_exp']
    E._table_from_df(doc, summary[cols], maxrows=40)
    doc.add_paragraph(
        'test_R2/RMSE/MSE и val_* — точность прогноза по NDVI на тесте и валидации. медиана_corr — '
        'медианная корреляция maxNDVI↔урожайность по районам культуры; доля_|r|>0.5 — доля районов с '
        'заметной связью; доля_exp — доля примеров, спрогнозированных экспонентой (а не откатом к среднему).')
    doc.add_heading('Сводный график', level=1)
    if bar:
        doc.add_picture(bar, width=Inches(6.3))

    doc.add_heading('Главы по культурам', level=1)
    for _, r in summary.iterrows():
        rus = r['культура']
        doc.add_heading(rus, level=2)
        doc.add_paragraph(
            f'Примеров: {int(r["n"])}, районов: {int(r["n_районов"])}. '
            f'ТЕСТ: R²={r["test_R2"]}, RMSE={r["test_RMSE"]}, MSE={r["test_MSE"]}. '
            f'ВАЛИДАЦИЯ: R²={r["val_R2"]}, RMSE={r["val_RMSE"]}, MSE={r["val_MSE"]}. '
            f'Корреляция maxNDVI↔урожай по районам: медиана {r["медиана_corr"]}, '
            f'|r|>0.5 у {r["доля_|r|>0.5"]*100:.0f}% районов. Доля прогнозов экспонентой: {r["доля_exp"]*100:.0f}%. '
            f'Полная таблица по районам — district_{[n for n,ru in NAMES if ru==rus][0]}.csv.')
        if sc_imgs.get(rus):
            doc.add_picture(sc_imgs[rus], width=Inches(4.2))

    doc.add_heading('Комментарии', level=1)
    for s in [
        'Сила связи NDVI↔урожай и точность сильно различаются по культурам: где вегетационный пик '
        'хорошо отражает урожай (зерновые), модель точнее; где урожай определяется не зелёной массой '
        '(корнеплоды, технические), связь слабее.',
        'Модель локальна (своя экспонента на каждый район) и использует один показатель — это быстрый '
        'интерпретируемый ориентир, но в целом уступает многомерным моделям (свёрточная сеть, ансамбль).',
        'Полные таблицы по районам (корреляции и коэффициенты a,b) — в district_<культура>.csv; '
        'сводка — summary_all_cultures.csv.']:
        doc.add_paragraph(s, style='List Bullet')

    out = os.path.join(docs, 'Регрессия_NDVI_по_культурам.docx')
    try:
        doc.save(out)
    except PermissionError:
        out = os.path.join(docs, 'Регрессия_NDVI_по_культурам_new.docx'); doc.save(out)
    print(f'Отчёт: {out}')


if __name__ == '__main__':
    main()
