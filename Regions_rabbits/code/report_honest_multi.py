"""ЧЕСТНАЯ многолетняя кривая заблаговременности (winter_wheat_honest_e2e).

В отличие от прежнего опыта (плоско ~0.85 из-за утечки t-1), здесь на КАЖДОЙ отсечке
вспомогательные признаки честные: prod_hist и климатология берутся только из данных,
известных на момент прогноза (год-прогноз b=a//size). Горизонт — ПОЛНЫЙ многолетний (0..~940 дн).

Запуск:  python code/report_honest_multi.py [--ver v3]
Выход:   workspace/reports/honest_multi_winter_wheat/
"""

import argparse
import json
import os

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch

import config
import explore_models as E
import plots
import train_honest as TH
import train_nn as T

GROUP = 'winter_wheat_honest_e2e_nn'
STEP_DAYS = 30


def _prep(dataset, seed):
    z, meta = TH.load_honest(dataset)
    raw, clim, ph, valid = z['raw'], z['clim'], z['ph'], z['valid']
    ids, y = z['ids'], z['y'].astype(np.float32)
    keep = ~np.isnan(y) & (y >= 0) & (y <= E.YIELD_MAX)
    raw, clim, ph, valid, ids, y = raw[keep], clim[keep], ph[keep], valid[keep], ids[keep], y[keep]
    n_years = int(meta['duration_years']); size = int(meta['time_rows_size']); P = int(meta['n_params'])
    tr, te, va = TH.split_random(len(y), seed)
    rawn, climn, phn, _ = TH.normalize(raw, clim, ph, tr)
    rawf = rawn.reshape(len(y), P, n_years * size)
    return dict(rawf=rawf, climn=climn, phn=phn, valid=valid, y=y, ids=ids,
                te=te, va=va, n_years=n_years, size=size, P=P, meta=meta)


def curves(dataset, ver, seed=42):
    mdir = os.path.join(config.MODELS_DIR, GROUP, ver)
    arch = json.load(open(os.path.join(mdir, 'architecture.json'), encoding='utf-8'))
    det = json.load(open(os.path.join(mdir, 'details.json'), encoding='utf-8'))
    D = _prep(dataset, int(det.get('seed', seed)))
    model = T.make_model(arch['arch_name'], arch['in_channels'], arch['in_length'])[0]
    model.load_state_dict(torch.load(os.path.join(mdir, 'model.pt'), map_location='cpu'))
    model.eval()

    n_years, size = D['n_years'], D['size']
    fill = det.get('fill', 'clim'); clim_ch = bool(det.get('clim_channels', True))
    harvest_abs = (n_years - 1) * size + int(det['harvest_doy'])
    leads = sorted(set(range(0, harvest_abs + 1, STEP_DAYS)) | {harvest_abs})
    out = {}
    for split in ('test', 'val'):
        idx = D['te'] if split == 'test' else D['va']
        rows = []
        for L in leads:
            a = max(0, harvest_abs - L)
            b = min(a // size, n_years - 1)
            sel = idx[D['valid'][idx, b]]                         # только примеры с климатологией года b
            if len(sel) < 20:
                continue
            X = TH.build_input(D['rawf'][sel], D['climn'][sel], D['phn'][sel], a, b, n_years, size,
                               fill=fill, clim_channels=clim_ch)
            pred = T.predict(model, X, 'cpu')
            rows.append({'days_before_harvest': L, 'n': len(sel), **E.metrics(D['y'][sel], pred)})
        out[split] = pd.DataFrame(rows)
    return out, D, det


def main(ver='v12', dataset='winter_wheat_honest'):
    outdir = os.path.join(config.REPORTS_DIR, 'honest_multi_winter_wheat')
    os.makedirs(outdir, exist_ok=True)
    res, D, det = curves(dataset, ver)

    def fig(metric):
        def draw(lang):
            fig, ax = plt.subplots(figsize=(10, 5.5))
            for split, c in res.items():
                lab = 'test (случ.)' if split == 'test' else 'val (случ.)'
                ax.plot(c.days_before_harvest, c[metric], marker='.', label=lab)
            ax.invert_xaxis()
            if metric == 'R2':
                ax.axhline(0, color='gray', lw=0.8, ls='--')
            ax.set_xlabel(plots.tr('Дней до уборки (время →, ближе к уборке)',
                                   'Days before harvest (time →, toward harvest)', lang))
            ax.set_ylabel(metric)
            ax.set_title(plots.tr(
                f'ЧЕСТНО (многолетне): заблаговременность {metric} — озимая пшеница',
                f'HONEST (multi-year): forecast lead time {metric} — winter wheat', lang))
            ax.grid(True, alpha=0.3); ax.legend(fontsize=9)
            fig.tight_layout()
            return fig
        return draw

    imgs = [('R² (val/test, случайный сплит)', plots.bilingual(outdir, 'leadtime_r2', fig('R2'))),
            ('MSE (val/test, случайный сплит)', plots.bilingual(outdir, 'leadtime_mse', fig('MSE')))]

    for split, c in res.items():
        c.to_csv(os.path.join(outdir, f'leadtime_{split}.csv'), index=False, encoding='utf-8-sig')

    _docx(outdir, res, det, imgs)
    print('\n=== ВАЛИДАЦИЯ (случайный сплит, честно многолетне) ===')
    print(res['val'][['days_before_harvest', 'n', 'R2', 'MSE']].to_string(index=False))
    print('\n=== ТЕСТ (случайный сплит) ===')
    print(res['test'][['days_before_harvest', 'n', 'R2', 'MSE']].to_string(index=False))
    print(f'\nГотово: {outdir}')


def _docx(outdir, res, det, imgs):
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading('ЧЕСТНАЯ многолетняя заблаговременность — озимая пшеница', 0)
    doc.add_paragraph(
        'Вспомогательные признаки (средняя урожайность и климатология) рассчитаны честно — по данным, '
        'доступным на момент прогноза: на отсечке a год-прогноз b = a // 365; prod_hist = средняя '
        'урожайность годов <= (year_b - 1) (статистика за прошлый год известна к Новому году); '
        'климатология = среднее за годы до year_b. Это убирает прежнюю утечку (год t-1), из-за которой '
        'кривая ошибочно держалась ~0.85. Сплит train/test/val — СЛУЧАЙНЫЙ 70/15/15 по примерам. '
        'Горизонт — полный многолетний (0..~940 дней до уборки).')
    doc.add_paragraph(
        'Кривая показывает, как падает качество с ростом заблаговременности, когда вспомогательные '
        'признаки честные (не содержат будущего относительно даты прогноза). Ступени на границах лет '
        'отражают появление новой статистики урожайности к Новому году.')

    for split in ('val', 'test'):
        lab = 'test (случ.)' if split == 'test' else 'val (случ.)'
        doc.add_heading(f'{lab} — по горизонтам', level=1)
        t = res[split].copy()
        for col in ('R2', 'RMSE', 'MSE', 'R_pearson'):
            if col in t:
                t[col] = t[col].round(3)
        E._table_from_df(doc, t[['days_before_harvest', 'n', 'R2', 'RMSE', 'MSE', 'R_pearson']], maxrows=60)

    doc.add_heading('Графики', level=1)
    for title, png in imgs:
        doc.add_heading(title, level=2)
        if png:
            doc.add_picture(png, width=Inches(6.3))

    out = os.path.join(outdir, 'report_honest_multi_winter_wheat.docx')
    try:
        doc.save(out)
    except PermissionError:
        out = os.path.join(outdir, 'report_honest_multi_winter_wheat_new.docx'); doc.save(out)
    print(f'Отчёт: {out}')


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--ver', default='v12')
    ap.add_argument('--dataset', default='winter_wheat_honest')
    a = ap.parse_args()
    main(a.ver, a.dataset)
