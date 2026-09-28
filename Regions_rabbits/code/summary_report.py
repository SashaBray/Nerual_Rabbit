"""Сводный отчёт по всем экспериментам: сначала РАЗВЕДКА (классические модели), затем
обученные НЕЙРОСЕТИ. Таблицы, графики, комментарии и сведения о размерах датасетов.

Источники:
  * размеры датасетов  — ``workspace/datasets/<task>/meta.json``
  * разведка           — ``workspace/reports/experiments/<task>/metrics.csv``
  * нейросети          — ``workspace/reports/nn/**/metrics.csv`` (+ детали из
                         зеркальной ``workspace/models/**/details.json|architecture.json``)

Выход: ``workspace/reports/summary/`` — ``summary_report.docx`` + ``datasets.csv`` +
``eda_summary.csv`` + ``nn_summary.csv`` + графики ``*.png``.

Запуск:  python code/summary_report.py
"""

import json
import os

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

import config
import explore_models as E   # переиспользуем _table_from_df
import plots


# --------------------------------------------------------------------------- #
# Сбор данных
# --------------------------------------------------------------------------- #
def _meta(task):
    p = os.path.join(config.dataset_dir(task), 'meta.json')
    return json.load(open(p, encoding='utf-8')) if os.path.exists(p) else {}


def collect_datasets():
    rows = []
    for task in sorted(os.listdir(config.DATASETS_DIR)):
        m = _meta(task)
        if not m:
            continue
        yrs = m.get('years', [])
        rows.append({
            'dataset': task, 'culture': m.get('culture_title', ''),
            'n_samples': m.get('n_samples'), 'features': m.get('n_features'),
            'duration_years': m.get('duration_years'),
            'time_len': m.get('time_rows_size', 0) * m.get('duration_years', 0),
            'matrix_width': m.get('matrix_width'), 'prod_hist_last': m.get('prod_hist_last'),
            'years': f'{min(yrs)}-{max(yrs)}' if yrs else '',
        })
    return pd.DataFrame(rows)


def collect_eda():
    exp = os.path.join(config.REPORTS_DIR, 'experiments')
    rows = []
    if os.path.isdir(exp):
        for task in sorted(os.listdir(exp)):
            mp = os.path.join(exp, task, 'metrics.csv')
            if not os.path.exists(mp):
                continue
            m = pd.read_csv(mp)
            meta = _meta(task)
            rec = {'dataset': task, 'culture': meta.get('culture_title', ''),
                   'n_samples': meta.get('n_samples')}
            for model, pref in (('baseline_prodhist', 'base'), ('linear', 'linear'),
                                ('tree', 'tree'), ('forest', 'forest')):
                sub = m[(m.model == model) & (m.split == 'test')]
                if len(sub):
                    rec[f'{pref}_R2'] = round(float(sub.iloc[0].R2), 3)
                    rec[f'{pref}_RMSE'] = round(float(sub.iloc[0].RMSE), 2)
            rows.append(rec)
    df = pd.DataFrame(rows)
    if 'forest_R2' in df:
        df = df.sort_values('forest_R2', ascending=False).reset_index(drop=True)
    return df


def collect_nn():
    """Идём от РЕАЛЬНО существующих моделей в models/ (с model.pt); метрики берём из
    зеркального отчёта reports/nn/<rel>/metrics.csv. Осиротевшие отчёты (модель удалена) игнорируем."""
    models = config.MODELS_DIR
    rows = []
    if os.path.isdir(models):
        for root, _dirs, files in os.walk(models):
            if 'model.pt' not in files:
                continue
            rel = os.path.relpath(root, models).replace('\\', '/')   # group/vN
            mp = os.path.join(config.REPORTS_DIR, 'nn', *rel.split('/'), 'metrics.csv')
            if not os.path.exists(mp):
                continue
            m = pd.read_csv(mp)
            det = _load_json(os.path.join(root, 'details.json'))
            arch = _load_json(os.path.join(root, 'architecture.json'))
            task = det.get('task', '')
            meta = _meta(task)

            def g(split, col):
                sub = m[m.split == split]
                return round(float(sub.iloc[0][col]), 3) if len(sub) else None

            rows.append({
                'model': rel, 'culture': meta.get('culture_title', ''), 'task': task,
                'arch': det.get('arch', arch.get('arch_name', '')),
                'lr': det.get('lr'), 'batch': det.get('batch'),
                'epochs': det.get('epochs_run'), 'best_epoch': det.get('best_epoch'),
                'n_params': arch.get('n_params'),
                'test_R2': g('test', 'R2'), 'test_RMSE': g('test', 'RMSE'), 'test_MSE': g('test', 'MSE'),
                'val_R2': g('val', 'R2'), 'val_RMSE': g('val', 'RMSE'),
                'n_samples': meta.get('n_samples'),
            })
    df = pd.DataFrame(rows)
    if 'test_R2' in df:
        df = df.sort_values('test_R2', ascending=False).reset_index(drop=True)
    return df


def _load_json(p):
    try:
        return json.load(open(p, encoding='utf-8'))
    except Exception:                       # noqa: BLE001
        return {}


# --------------------------------------------------------------------------- #
# Графики
# --------------------------------------------------------------------------- #
def _barplot(labels, values, outdir, name, title_ru, title_en, color='tab:blue'):
    labels = list(labels)

    def draw(lang):
        fig, ax = plt.subplots(figsize=(11, 5))
        ax.bar(range(len(labels)), values, color=color)
        ax.set_xticks(range(len(labels)))
        ax.set_xticklabels(labels, rotation=60, ha='right', fontsize=8)
        ax.set_ylabel('test R²'); ax.set_title(plots.tr(title_ru, title_en, lang))
        ax.grid(True, axis='y', alpha=0.3)
        fig.tight_layout()
        return fig
    return plots.bilingual(outdir, name, draw)


def _compare_plot(merged, outdir):
    x = np.arange(len(merged)); w = 0.4

    def draw(lang):
        fig, ax = plt.subplots(figsize=(11, 5))
        ax.bar(x - w / 2, merged.forest_R2, w,
               label=plots.tr('RandomForest (разведка)', 'RandomForest (exploration)', lang))
        ax.bar(x + w / 2, merged.nn_R2, w,
               label=plots.tr('Нейросеть (лучшая)', 'Neural net (best)', lang))
        ax.set_xticks(x); ax.set_xticklabels(merged.dataset, rotation=60, ha='right', fontsize=8)
        ax.set_ylabel('test R²')
        ax.set_title(plots.tr('Классика vs нейросеть по культурам',
                              'Classical vs neural net by culture', lang))
        ax.legend(); ax.grid(True, axis='y', alpha=0.3)
        fig.tight_layout()
        return fig
    return plots.bilingual(outdir, 'compare', draw)


# --------------------------------------------------------------------------- #
# Комментарии (авто)
# --------------------------------------------------------------------------- #
def comments_datasets(ds):
    c = [f'Всего датасетов: {len(ds)}. Размеры (n_samples) от {int(ds.n_samples.min())} '
         f'до {int(ds.n_samples.max())}, суммарно {int(ds.n_samples.sum())} примеров.']
    big = ds.sort_values('n_samples', ascending=False).iloc[0]
    small = ds.sort_values('n_samples').iloc[0]
    c.append(f'Самый крупный: {big.culture or big.dataset} ({int(big.n_samples)}); '
             f'самый мелкий: {small.culture or small.dataset} ({int(small.n_samples)}) — '
             f'мелкие дают шумные метрики.')
    c.append(f'У всех {int(ds.features.mode().iloc[0])} признаков, длина ряда '
             f'{int(ds.time_len.mode().iloc[0])} точек (+1 строка скользящего среднего урожайности).')
    return c


def comments_eda(eda):
    if not len(eda) or 'forest_R2' not in eda:
        return ['Разведочных отчётов не найдено.']
    c = [f'Разведка проведена по {len(eda)} датасетам (модели: линейная, дерево, случайный лес).',
         f'Средний R² случайного леса (test): {eda.forest_R2.mean():.2f}.']
    best, worst = eda.iloc[0], eda.iloc[-1]
    c.append(f'Лучше всего предсказывается {best.culture or best.dataset} '
             f'(лес R²={best.forest_R2}), хуже всего — {worst.culture or worst.dataset} '
             f'(R²={worst.forest_R2}).')
    if 'linear_R2' in eda:
        d = (eda.forest_R2 - eda.linear_R2).mean()
        c.append(f'Случайный лес в среднем сильнее линейной модели на {d:.2f} по R² — '
                 f'связь признаков с урожайностью заметно нелинейная.')
    return c


def comments_nn(nn):
    if not len(nn):
        return ['Обученных нейросетей пока нет (запустите train_nn.py).']
    c = [f'Обучено моделей/версий: {len(nn)} по {nn.task.nunique()} культурам.']
    best = nn.iloc[0]
    c.append(f'Лучшая сеть: {best.model} ({best.culture}) — test R²={best.test_R2}, '
             f'RMSE={best.test_RMSE}, MSE={best.test_MSE}.')
    if nn['arch'].nunique() == 1:
        c.append(f'Архитектура: {best.arch}; параметров ~{int(best.n_params):,}.')
    return c


def comments_compare(merged):
    if not len(merged):
        return []
    win = int((merged.nn_R2 > merged.forest_R2).sum())
    c = [f'Сравнение есть по {len(merged)} культурам (где обучена сеть). '
         f'Нейросеть обошла случайный лес в {win} из {len(merged)}.']
    diff = (merged.nn_R2 - merged.forest_R2)
    c.append(f'Средняя разница R² (сеть − лес): {diff.mean():+.3f}.')
    return c


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #
def _disp(df):
    """Копия для таблицы: NaN/None -> '—' (чтобы в docx не было 'None')."""
    return df.fillna('—')


def build_docx(ds, eda, nn, imgs, cm, outdir):
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading('Сводный отчёт: разведка и нейросети', 0)

    doc.add_heading('1. Датасеты (размеры)', 1)
    for s in cm['datasets']:
        doc.add_paragraph(s)
    E._table_from_df(doc, _disp(ds))

    doc.add_heading('2. Разведка — классические модели (test)', 1)
    for s in cm['eda']:
        doc.add_paragraph(s)
    if len(eda):
        E._table_from_df(doc, _disp(eda))
    if imgs.get('eda_bar'):
        doc.add_picture(imgs['eda_bar'], width=Inches(6.5))

    doc.add_heading('3. Нейросети (test)', 1)
    for s in cm['nn']:
        doc.add_paragraph(s)
    if len(nn):
        E._table_from_df(doc, _disp(nn))
    if imgs.get('nn_bar'):
        doc.add_picture(imgs['nn_bar'], width=Inches(6.5))

    if imgs.get('compare'):
        doc.add_heading('4. Классика vs нейросеть', 1)
        for s in cm['compare']:
            doc.add_paragraph(s)
        doc.add_picture(imgs['compare'], width=Inches(6.5))

    p = os.path.join(outdir, 'summary_report.docx')
    doc.save(p)
    return p


# --------------------------------------------------------------------------- #
# Главный сценарий
# --------------------------------------------------------------------------- #
def run():
    outdir = os.path.join(config.REPORTS_DIR, 'summary')
    os.makedirs(outdir, exist_ok=True)

    ds = collect_datasets()
    eda = collect_eda()
    nn = collect_nn()
    ds.to_csv(os.path.join(outdir, 'datasets.csv'), index=False, encoding='utf-8-sig')
    eda.to_csv(os.path.join(outdir, 'eda_summary.csv'), index=False, encoding='utf-8-sig')
    nn.to_csv(os.path.join(outdir, 'nn_summary.csv'), index=False, encoding='utf-8-sig')

    imgs = {}
    if len(eda) and 'forest_R2' in eda:
        imgs['eda_bar'] = _barplot(eda.dataset, eda.forest_R2, outdir, 'eda_forest_R2',
                                   'Разведка: R² случайного леса (test) по культурам',
                                   'Exploration: RandomForest R² (test) by culture')
    merged = pd.DataFrame()
    if len(nn) and 'test_R2' in nn:
        best_nn = nn.groupby('task', as_index=False).agg(nn_R2=('test_R2', 'max'),
                                                         culture=('culture', 'first'))
        imgs['nn_bar'] = _barplot(best_nn.culture.where(best_nn.culture != '', best_nn.task),
                                  best_nn.nn_R2, outdir, 'nn_best_R2',
                                  'Нейросети: лучший R² (test) по культурам',
                                  'Neural nets: best R² (test) by culture', color='tab:green')
        if len(eda):
            merged = eda[['dataset', 'forest_R2']].merge(
                best_nn, left_on='dataset', right_on='task', how='inner')
            merged = merged.sort_values('forest_R2', ascending=False).reset_index(drop=True)
            if len(merged):
                imgs['compare'] = _compare_plot(merged, outdir)

    cm = {'datasets': comments_datasets(ds), 'eda': comments_eda(eda),
          'nn': comments_nn(nn), 'compare': comments_compare(merged)}
    rep = build_docx(ds, eda, nn, imgs, cm, outdir)

    print(f'Датасетов: {len(ds)}, разведочных отчётов: {len(eda)}, нейросетей: {len(nn)}')
    print(f'Сводный отчёт: {rep}')
    print(f'Таблицы: {outdir} (datasets.csv, eda_summary.csv, nn_summary.csv)')


if __name__ == '__main__':
    run()
