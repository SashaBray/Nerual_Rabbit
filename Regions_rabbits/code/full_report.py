"""Единый отчёт обо ВСЕХ экспериментах: обзор + ГЛАВА НА КАЖДУЮ модель (культуру).

Каждая глава: датасет (размеры), разведка (классические модели + свип L + важность),
нейросеть (конфигурация + метрики + кривая обучения), заблаговременный прогноз (таблица + графики).
В конце — сводные таблицы и сравнительные графики.

Использует уже сгенерированные артефакты:
  reports/experiments/<task>/, models/<task>_nn/vN/, reports/nn/<task>_nn/vN/,
  reports/forecast/<task>_nn_v*/, reports/summary/ (создаётся здесь через summary_report).

Запуск:  python code/full_report.py   ->  workspace/reports/full_report.docx
"""

import json
import os

import pandas as pd

import config
import explore_models as E
import summary_report as S


def _j(p):
    try:
        return json.load(open(p, encoding='utf-8'))
    except Exception:                       # noqa: BLE001
        return {}


def _png(*parts):
    p = os.path.join(*parts)
    return p if os.path.exists(p) else None


def _latest_ver(group_dir):
    if not os.path.isdir(group_dir):
        return None
    vs = [d for d in os.listdir(group_dir) if d.startswith('v') and d[1:].isdigit()]
    return max(vs, key=lambda d: int(d[1:])) if vs else None


def _pic(doc, png, inches=6.0):
    from docx.shared import Inches
    if png:
        doc.add_picture(png, width=Inches(inches))


def _tasks_sorted():
    """Культуры, отсортированные по test-R² нейросети (лучшие — первыми)."""
    nn = S.collect_nn()
    order = {}
    if len(nn):
        best = nn.groupby('task')['test_R2'].max()
        order = best.to_dict()
    tasks = [d for d in sorted(os.listdir(config.DATASETS_DIR))
             if os.path.exists(os.path.join(config.dataset_dir(d), 'meta.json'))]
    return sorted(tasks, key=lambda t: -order.get(t, -1))


def chapter(doc, task):
    from docx.shared import Pt
    meta = _j(os.path.join(config.dataset_dir(task), 'meta.json'))
    culture = meta.get('culture_title', task)
    doc.add_heading(f'{culture} ({task})', level=1)

    # --- датасет ---
    doc.add_heading('Датасет', level=2)
    yrs = meta.get('years', [])
    doc.add_paragraph(
        f'Примеров: {meta.get("n_samples")}; признаков: {meta.get("n_features")} '
        f'(+1 строка скользящего среднего урожайности); длина ряда '
        f'{meta.get("time_rows_size")}×{meta.get("duration_years")} точек; '
        f'годы {min(yrs) if yrs else "?"}–{max(yrs) if yrs else "?"}; '
        f'окно скольз. среднего урожайности m={meta.get("prod_hist_last")}.')

    # --- разведка ---
    doc.add_heading('Разведка: классические модели', level=2)
    edir = os.path.join(config.REPORTS_DIR, 'experiments', task)
    mp = os.path.join(edir, 'metrics.csv')
    if os.path.exists(mp):
        E._table_from_df(doc, pd.read_csv(mp).round(3))
        imp = os.path.join(edir, 'importance_forest.csv')
        if os.path.exists(imp):
            doc.add_paragraph('Важность признаков (случайный лес), топ:')
            E._table_from_df(doc, pd.read_csv(imp).head(6).round(3))
        _pic(doc, _png(edir, 'lengths_sweep_ru.png'))
        _pic(doc, _png(edir, 'scatter_forest_ru.png'), 4.5)
    else:
        doc.add_paragraph('— нет отчёта разведки.')

    # --- нейросеть ---
    doc.add_heading('Нейросеть', level=2)
    gdir = os.path.join(config.MODELS_DIR, f'{task}_nn')
    ver = _latest_ver(gdir)
    if ver:
        det = _j(os.path.join(gdir, ver, 'details.json'))
        arch = _j(os.path.join(gdir, ver, 'architecture.json'))
        doc.add_paragraph(
            f'Версия {ver}; архитектура {det.get("arch")}; параметров '
            f'~{arch.get("n_params", "?"):,}; lr={det.get("lr")}, batch={det.get("batch")}, '
            f'эпох {det.get("epochs_run")} (лучшая {det.get("best_epoch")}).')
        rdir = os.path.join(config.REPORTS_DIR, 'nn', f'{task}_nn', ver)
        mp = os.path.join(rdir, 'metrics.csv')
        if os.path.exists(mp):
            E._table_from_df(doc, pd.read_csv(mp).round(3))
        _pic(doc, _png(rdir, 'training_curve_ru.png'))
        _pic(doc, _png(rdir, 'scatter_test_ru.png'), 4.5)
    else:
        doc.add_paragraph('— модель не обучена.')

    # --- заблаговременность ---
    doc.add_heading('Заблаговременный прогноз', level=2)
    fdir = None
    base = os.path.join(config.REPORTS_DIR, 'forecast')
    if ver and os.path.isdir(os.path.join(base, f'{task}_nn_{ver}')):
        fdir = os.path.join(base, f'{task}_nn_{ver}')
    if fdir and os.path.exists(os.path.join(fdir, 'leadtime.csv')):
        df = pd.read_csv(os.path.join(fdir, 'leadtime.csv'))
        cols = [c for c in ['days_before_harvest', 'weeks_known', 'MSE', 'RMSE', 'R2'] if c in df.columns]
        E._table_from_df(doc, df[cols].round(3))
        _pic(doc, _png(fdir, 'leadtime_rmse_ru.png'))
        _pic(doc, _png(fdir, 'leadtime_mse_ru.png'))
    else:
        doc.add_paragraph('— нет эксперимента заблаговременности.')

    doc.add_page_break()


def build():
    from docx import Document
    from docx.shared import Inches
    S.run()                                  # сгенерировать сводные таблицы/графики

    doc = Document()
    doc.add_heading('Прогноз урожайности: полный отчёт по экспериментам', 0)
    doc.add_paragraph('Пайплайн: датасеты (хронологический порядок годов) → разведка (классические '
                      'модели со свипом длины сжатия) → свёрточная нейросеть (mse24) → эксперимент '
                      'заблаговременного прогноза относительно даты уборки. Все метрики — на ТЕСТОВОЙ '
                      'выборке; разбиение 70/15/15 фиксировано сидом.')

    # обзор
    sdir = os.path.join(config.REPORTS_DIR, 'summary')
    doc.add_heading('Обзор: размеры датасетов', level=1)
    if os.path.exists(os.path.join(sdir, 'datasets.csv')):
        E._table_from_df(doc, pd.read_csv(os.path.join(sdir, 'datasets.csv')).fillna('—'))
    doc.add_heading('Обзор: разведка vs нейросеть', level=1)
    for f in ('eda_summary.csv', 'nn_summary.csv'):
        if os.path.exists(os.path.join(sdir, f)):
            E._table_from_df(doc, pd.read_csv(os.path.join(sdir, f)).fillna('—'))
    _pic(doc, _png(sdir, 'compare_ru.png'), 6.5)
    fsum = os.path.join(config.REPORTS_DIR, 'forecast', '_summary')
    doc.add_heading('Обзор: заблаговременность по культурам', level=1)
    _pic(doc, _png(fsum, 'leadtime_all_ru.png'), 6.5)
    _pic(doc, _png(fsum, 'leadtime_all_mse_ru.png'), 6.5)
    doc.add_page_break()

    doc.add_heading('Главы по моделям', level=0)
    for task in _tasks_sorted():
        chapter(doc, task)

    out = os.path.join(config.REPORTS_DIR, 'full_report.docx')
    try:
        doc.save(out)
    except PermissionError:
        out = os.path.join(config.REPORTS_DIR, 'full_report_new.docx')
        doc.save(out)
        print('! full_report.docx занят (открыт в Word?) — сохранил как full_report_new.docx')
    print(f'Отчёт: {out}')


if __name__ == '__main__':
    build()
