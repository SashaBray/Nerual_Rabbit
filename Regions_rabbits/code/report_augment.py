"""Сравнение версий озимой пшеницы по заблаговременности — на ТЕСТЕ и ВАЛИДАЦИИ.

Версии группы winter_wheat_all_nn:
  v1 — без historical, без аугментации        v2 — historical, без аугментации
  v3 — historical + аугментация                v4 — без historical, с аугментацией
Для каждой версии строит кривую заблаговременности (R² от «дней до уборки») отдельно на test и val,
сравнительные графики (RU/EN, png+pdf) и Word-отчёт.

Запуск:  python code/report_augment.py
Выход:   workspace/reports/augment_winter_wheat/
"""

import json
import os

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch

import config
import explore_models as E
import forecast_leadtime as FC
import plots
import train_nn as T

GROUP = 'winter_wheat_all_nn'
STEP_DAYS = 30
HORIZONS = [0, 90, 180, 365, 545, 730, 900]
_PREP = {}


def _prep(task, seed):
    if task not in _PREP:
        P = E.prepare_data(task, seed)
        _PREP[task] = {'Xn': P['Xn'].astype(np.float32), 'y': P['y'].astype(np.float32),
                       'test': P['te'], 'val': P['va'], 'meta': P['meta']}
    return _PREP[task]


def _label(ver, in_ch, aug):
    return f'{ver} ({"hist" if in_ch >= 20 else "без hist"}, {"aug" if aug else "без aug"})'


def _curves(ver):
    """Вернуть {'label':..., 'test':df, 'val':df} или None, если модель несовместима с датасетом."""
    mdir = os.path.join(config.MODELS_DIR, GROUP, ver)
    if not os.path.exists(os.path.join(mdir, 'model.pt')):
        return None
    arch = json.load(open(os.path.join(mdir, 'architecture.json'), encoding='utf-8'))
    det = json.load(open(os.path.join(mdir, 'details.json'), encoding='utf-8'))
    task, seed, size = det['task'], int(det.get('seed', 42)), int(det['time_rows_size'])
    P = _prep(task, seed)
    Xn, y, meta = P['Xn'], P['y'], P['meta']
    F, Tlen = Xn.shape[1], Xn.shape[2]
    if int(arch['in_channels']) != F:
        print(f'  {ver}: каналы модели {arch["in_channels"]} != данных {F} (датасет пересобран) — пропуск')
        return None

    model = T.make_model(arch['arch_name'], arch['in_channels'], arch['in_length'])[0]
    model.load_state_dict(torch.load(os.path.join(mdir, 'model.pt'), map_location='cpu'))
    model.eval()

    n_years = Tlen // size
    tgt_idx = 0 if bool(meta.get('concat_newest_first', True)) else n_years - 1
    blocks = Xn.reshape(Xn.shape[0], F, n_years, size)
    clim = np.tile(np.delete(blocks, tgt_idx, axis=2).mean(axis=2), (1, 1, n_years)).astype(np.float32)
    harvest_abs = tgt_idx * size + FC._harvest_doy(task)
    leads = sorted(set(range(0, harvest_abs + 1, STEP_DAYS)) | {harvest_abs})

    out = {'label': _label(ver, F, det.get('augment_leadtime'))}
    for split in ('test', 'val'):
        idx = P[split]
        Xi, yi, ci = Xn[idx], y[idx], clim[idx]
        rows = []
        for lead_d in leads:
            a = max(0, min(Tlen, harvest_abs - lead_d))
            Xm = Xi.copy()
            if a < Tlen:
                Xm[:, :, a:] = ci[:, :, a:]
            m = E.metrics(yi, T.predict(model, Xm, 'cpu'))
            rows.append({'days_before_harvest': lead_d, **m})
        out[split] = pd.DataFrame(rows).sort_values('days_before_harvest').reset_index(drop=True)
    print(f'  {ver}: посчитано (test+val), {out["label"]}')
    return out


def _at(df, d, col='R2'):
    return float(df.loc[(df.days_before_harvest - d).abs().idxmin(), col])


def main():
    outdir = os.path.join(config.REPORTS_DIR, 'augment_winter_wheat')
    os.makedirs(outdir, exist_ok=True)
    res = {}
    for ver in ('v1', 'v2', 'v3', 'v4'):
        c = _curves(ver)
        if c:
            res[ver] = c
    if not res:
        raise SystemExit('Нет совместимых версий.')

    # графики: R² и MSE от заблаговременности, отдельно на val и test
    def fig_split(split, metric):
        def draw(lang):
            fig, ax = plt.subplots(figsize=(10, 5.5))
            for ver, c in res.items():
                d = c[split]
                ax.plot(d.days_before_harvest, d[metric], marker='.', label=c['label'])
            ax.invert_xaxis()
            ax.set_xlabel(plots.tr('Дней до уборки (время →, ближе к уборке)',
                                   'Days before harvest (time →, toward harvest)', lang))
            ax.set_ylabel(f'{metric} ({split})')
            ax.set_title(plots.tr(f'Заблаговременность ({metric}, {split}) — озимая пшеница',
                                  f'Forecast lead time ({metric}, {split}) — winter wheat', lang))
            ax.grid(True, alpha=0.3); ax.legend(fontsize=8)
            fig.tight_layout()
            return fig
        return draw
    imgs = [
        ('Валидация — R²', plots.bilingual(outdir, 'leadtime_val', fig_split('val', 'R2'))),
        ('Валидация — MSE', plots.bilingual(outdir, 'leadtime_val_mse', fig_split('val', 'MSE'))),
        ('Тест — R²', plots.bilingual(outdir, 'leadtime_test', fig_split('test', 'R2'))),
        ('Тест — MSE', plots.bilingual(outdir, 'leadtime_test_mse', fig_split('test', 'MSE'))),
    ]

    # таблицы: R² по горизонтам на val и test
    def table(split):
        rows = []
        for ver, c in res.items():
            rec = {'версия': c['label']}
            for h in HORIZONS:
                rec[f'd{h}'] = round(_at(c[split], h), 3)
            rec['среднее'] = round(c[split].R2.mean(), 3)
            rows.append(rec)
        return pd.DataFrame(rows)
    tab_val, tab_test = table('val'), table('test')
    tab_val.to_csv(os.path.join(outdir, 'leadtime_val.csv'), index=False, encoding='utf-8-sig')
    tab_test.to_csv(os.path.join(outdir, 'leadtime_test.csv'), index=False, encoding='utf-8-sig')

    _docx(outdir, tab_val, tab_test, imgs)
    print('\n=== R² по горизонтам (ВАЛИДАЦИЯ) ===')
    print(tab_val.to_string(index=False))
    print(f'\nГотово: {outdir}')


def _docx(outdir, tab_val, tab_test, imgs):
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading('Аугментация заблаговременности — озимая пшеница (test + val)', 0)
    doc.add_paragraph(
        'Сравнение версий: v1 (без historical, без аугментации), v2 (historical, без аугментации), '
        'v3 (historical + аугментация), v4 (без historical, с аугментацией). Аугментация: каждый '
        'пример обучается с разными отсечками (шаг 7 дн.), будущее заменяется климатологией — модель '
        'учится прогнозировать при любой заблаговременности. Метрики на ТЕСТЕ и ВАЛИДАЦИИ '
        '(по столбцам — R² за N дней до уборки; «среднее» — по всем горизонтам).')

    doc.add_heading('Таблица R² — валидация', level=1)
    E._table_from_df(doc, tab_val)
    doc.add_heading('Таблица R² — тест', level=1)
    E._table_from_df(doc, tab_test)

    doc.add_heading('Графики (R² и MSE от дней до уборки)', level=1)
    for title, png in imgs:
        doc.add_heading(title, level=2)
        if png:
            doc.add_picture(png, width=Inches(6.3))

    out = os.path.join(outdir, 'report_augment_winter_wheat.docx')
    try:
        doc.save(out)
    except PermissionError:
        out = os.path.join(outdir, 'report_augment_winter_wheat_new.docx')
        doc.save(out)
        print('! исходный docx занят — сохранил как *_new.docx')
    print(f'Отчёт: {out}')


if __name__ == '__main__':
    main()
