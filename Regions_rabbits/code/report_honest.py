"""ЧЕСТНЫЙ эксперимент по заблаговременности — озимая пшеница.

Два исправления относительно прежнего опыта (где кривая ошибочно была плоской на ~0.85):
  1) СПЛИТ ПО ГОДАМ (leave-one-year-out): train<=2020, val=2021, test=2022 — проверяем
     настоящий прогноз будущего года, а не «память района» (раньше 98% тест-районов были и в train).
  2) ЧЕСТНЫЙ ДИАПАЗОН заблаговременности — только внутри целевого года (0..день уборки).
     В этом диапазоне prod_hist и климатология (годы t-1,t-2) — уже прошедшие, утечки t-1 нет.
     Аугментация тоже обучалась со scope=year (отсечки только в целевом году).

Версии группы winter_wheat_honest_nn:
  v1 — historical, без аугментации
  v2 — historical, с аугментацией (scope=year)
  v3 — без historical (raw), с аугментацией (scope=year)

Запуск:  python code/report_honest.py
Выход:   workspace/reports/honest_winter_wheat/
"""

import json
import os

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch

import config
import explore_models as E
import forecast_leadtime as FC
import plots
import train_nn as T

GROUP = 'winter_wheat_honest_nn'
STEP_DAYS = 15           # шаг внутри целевого года
_PREP = {}


def _prep(task, seed):
    if task not in _PREP:
        E.SPLIT_MODE = 'year'                      # ВАЖНО: те же индексы, что при обучении
        P = E.prepare_data(task, seed)
        _PREP[task] = {'Xn': P['Xn'].astype(np.float32), 'y': P['y'].astype(np.float32),
                       'test': P['te'], 'val': P['va'], 'meta': P['meta'],
                       'ids': P['ids'], 'split_years': P['split_years']}
    return _PREP[task]


def _label(det):
    hist = 'historical' if int(det['in_channels']) >= 20 else 'без historical (raw)'
    aug = f'aug scope={det.get("aug_scope")}' if det.get('augment_leadtime') else 'без аугментации'
    return f'{det["version"] if "version" in det else ""} ({hist}, {aug})'.strip()


def _curves(ver):
    mdir = os.path.join(config.MODELS_DIR, GROUP, ver)
    if not os.path.exists(os.path.join(mdir, 'model.pt')):
        return None
    arch = json.load(open(os.path.join(mdir, 'architecture.json'), encoding='utf-8'))
    det = json.load(open(os.path.join(mdir, 'details.json'), encoding='utf-8'))
    det['version'] = ver
    task, seed, size = det['task'], int(det.get('seed', 42)), int(det['time_rows_size'])
    P = _prep(task, seed)
    Xn, y, meta = P['Xn'], P['y'], P['meta']
    F, Tlen = Xn.shape[1], Xn.shape[2]
    if int(arch['in_channels']) != F:
        print(f'  {ver}: каналы модели {arch["in_channels"]} != данных {F} — пропуск')
        return None

    model = T.make_model(arch['arch_name'], arch['in_channels'], arch['in_length'])[0]
    model.load_state_dict(torch.load(os.path.join(mdir, 'model.pt'), map_location='cpu'))
    model.eval()

    n_years = Tlen // size
    tgt_idx = 0 if bool(meta.get('concat_newest_first', True)) else n_years - 1
    blocks = Xn.reshape(Xn.shape[0], F, n_years, size)
    clim = np.tile(np.delete(blocks, tgt_idx, axis=2).mean(axis=2), (1, 1, n_years)).astype(np.float32)
    H = FC._harvest_doy(task)
    harvest_abs = tgt_idx * size + H
    # ЧЕСТНЫЙ диапазон: только внутри целевого года -> lead 0..H (a остаётся в целевом блоке)
    leads = sorted(set(range(0, H + 1, STEP_DAYS)) | {H})

    out = {'label': _label(det), 'in_channels': F, 'aug': bool(det.get('augment_leadtime'))}
    for split in ('test', 'val'):
        idx = P[split]
        Xi, yi, ci = Xn[idx], y[idx], clim[idx]
        rows = []
        for lead_d in leads:
            a = max(0, min(Tlen, harvest_abs - lead_d))
            Xm = Xi.copy()
            if a < Tlen:
                Xm[:, :, a:] = ci[:, :, a:]
            m = E.metrics(yi, T.predict(model, Xm, 'cpu'))
            rows.append({'days_before_harvest': lead_d, **m})
        out[split] = pd.DataFrame(rows).sort_values('days_before_harvest').reset_index(drop=True)
    print(f'  {ver}: посчитано (test+val), {out["label"]}')
    return out


def _at(df, d, col='R2'):
    return float(df.loc[(df.days_before_harvest - d).abs().idxmin(), col])


def _year_table(task, seed):
    P = _prep(task, seed)
    df = pd.DataFrame({'year': P['ids'].year.values, 'y': P['y']})
    g = df.groupby('year')['y'].agg(['count', 'mean', 'std']).reset_index()
    g.columns = ['год', 'примеров', 'средняя_урожайность', 'std']
    g['средняя_урожайность'] = g['средняя_урожайность'].round(2)
    g['std'] = g['std'].round(2)
    return g


def main():
    outdir = os.path.join(config.REPORTS_DIR, 'honest_winter_wheat')
    os.makedirs(outdir, exist_ok=True)
    res = {}
    for ver in ('v1', 'v2', 'v3'):
        c = _curves(ver)
        if c:
            res[ver] = c
    if not res:
        raise SystemExit('Нет совместимых версий honest.')

    H = FC._harvest_doy('winter_wheat_all')

    def fig_split(split, metric):
        def draw(lang):
            fig, ax = plt.subplots(figsize=(10, 5.5))
            for ver, c in res.items():
                d = c[split]
                ax.plot(d.days_before_harvest, d[metric], marker='.', label=c['label'])
            ax.invert_xaxis()
            ax.axhline(0 if metric == 'R2' else np.nan, color='gray', lw=0.8, ls='--')
            ax.set_xlabel(plots.tr('Дней до уборки (время →, ближе к уборке)',
                                   'Days before harvest (time →, toward harvest)', lang))
            ax.set_ylabel(f'{metric} ({split})')
            yr = '2022' if split == 'test' else '2021'
            ax.set_title(plots.tr(
                f'ЧЕСТНО: заблаговременность ({metric}, {split}={yr}) — озимая пшеница',
                f'HONEST: forecast lead time ({metric}, {split}={yr}) — winter wheat', lang))
            ax.grid(True, alpha=0.3); ax.legend(fontsize=8)
            fig.tight_layout()
            return fig
        return draw

    imgs = [
        ('Валидация 2021 — R²', plots.bilingual(outdir, 'leadtime_val', fig_split('val', 'R2'))),
        ('Валидация 2021 — MSE', plots.bilingual(outdir, 'leadtime_val_mse', fig_split('val', 'MSE'))),
        ('Тест 2022 — R²', plots.bilingual(outdir, 'leadtime_test', fig_split('test', 'R2'))),
        ('Тест 2022 — MSE', plots.bilingual(outdir, 'leadtime_test_mse', fig_split('test', 'MSE'))),
    ]

    horizons = sorted(set(range(0, H + 1, 30)) | {H})

    def table(split):
        rows = []
        for ver, c in res.items():
            rec = {'версия': c['label']}
            for h in horizons:
                rec[f'd{h}'] = round(_at(c[split], h), 3)
            rec['среднее'] = round(c[split].R2.mean(), 3)
            rows.append(rec)
        return pd.DataFrame(rows)

    tab_val, tab_test = table('val'), table('test')
    tab_val.to_csv(os.path.join(outdir, 'leadtime_val.csv'), index=False, encoding='utf-8-sig')
    tab_test.to_csv(os.path.join(outdir, 'leadtime_test.csv'), index=False, encoding='utf-8-sig')
    year_tab = _year_table('winter_wheat_all', 42)
    year_tab.to_csv(os.path.join(outdir, 'year_levels.csv'), index=False, encoding='utf-8-sig')

    _docx(outdir, tab_val, tab_test, year_tab, imgs)
    print('\n=== R² по горизонтам (ВАЛИДАЦИЯ 2021) ===')
    print(tab_val.to_string(index=False))
    print('\n=== R² по горизонтам (ТЕСТ 2022) ===')
    print(tab_test.to_string(index=False))
    print(f'\nГотово: {outdir}')


def _docx(outdir, tab_val, tab_test, year_tab, imgs):
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading('ЧЕСТНЫЙ опыт заблаговременности — озимая пшеница', 0)
    doc.add_paragraph(
        'Прежний опыт давал подозрительно плоскую кривую ~0.85 на всём горизонте. Причины были '
        'выявлены и устранены:')
    doc.add_paragraph(
        '1) Сплит по ПРИМЕРАМ давал утечку «памяти района»: 98% тест-районов присутствовали и в '
        'train (по другим годам). Здесь — сплит ПО ГОДАМ: train ≤2020, валидация = 2021, тест = 2022.',
        style='List Bullet')
    doc.add_paragraph(
        '2) Многолетний горизонт давал временную утечку: вспомогательные признаки (prod_hist и '
        'климатология) содержат год t-1, который при раннем прогнозе ещё в будущем. Здесь честный '
        'горизонт — только ВНУТРИ целевого года (0..день уборки), где t-1, t-2 уже прошли. '
        'Аугментация обучалась со scope=year (отсечки только в целевом году).',
        style='List Bullet')

    doc.add_heading('Главный результат', level=1)
    doc.add_paragraph(
        'На новом году прогноз АБСОЛЮТНОГО уровня урожайности труден: тестовый 2022 г. — рекордный '
        '(средняя 41.9 ц/га против 33.5 в train), и модель, обученная на 2015–2020, систематически '
        'недооценивает его уровень → R² на тесте отрицательный, ХОТЯ корреляция Пирсона ~0.78 '
        '(пространственный рейтинг районов улавливается). Даже наивный прогноз «средним трейна» даёт '
        'R²≈−0.39 на 2022. На валидации 2021 (ближе к train по уровню) R² ≈ 0.5–0.69. '
        'Вывод: прежние ~0.85 были артефактом утечки; честная задача «спрогнозировать новый год» '
        'намного сложнее, а вклад аугментации в честном диапазоне невелик.')

    doc.add_heading('Уровень урожайности по годам (сдвиг распределения)', level=1)
    E._table_from_df(doc, year_tab)

    doc.add_heading('R² по горизонтам — валидация 2021', level=1)
    E._table_from_df(doc, tab_val)
    doc.add_heading('R² по горизонтам — тест 2022', level=1)
    E._table_from_df(doc, tab_test)

    doc.add_heading('Графики (R² и MSE от дней до уборки, честный горизонт)', level=1)
    for title, png in imgs:
        doc.add_heading(title, level=2)
        if png:
            doc.add_picture(png, width=Inches(6.3))

    out = os.path.join(outdir, 'report_honest_winter_wheat.docx')
    try:
        doc.save(out)
    except PermissionError:
        out = os.path.join(outdir, 'report_honest_winter_wheat_new.docx')
        doc.save(out)
        print('! исходный docx занят — сохранил как *_new.docx')
    print(f'Отчёт: {out}')


if __name__ == '__main__':
    main()
